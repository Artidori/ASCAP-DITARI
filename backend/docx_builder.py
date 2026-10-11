import io
import re
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm

FONT = "Times New Roman"
BLOCK = {"div", "p", "li"}


def _set_font(run, size, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold or None
    run.italic = italic or None
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def _style_flags(el, bold, italic):
    style = (el.get("style") or "").replace(" ", "").lower()
    bold = bold or el.name in ("b", "strong") or "font-weight:bold" in style or "font-weight:700" in style
    italic = italic or el.name in ("i", "em") or "font-style:italic" in style
    return bold, italic


def _fill(cell_el, cell, size, center, img_loader=None):
    cell.paragraphs[0].text = ""
    state = {"p": cell.paragraphs[0]}

    def new_par():
        state["p"] = cell.add_paragraph()

    def walk(node, bold, italic):
        for ch in node.children:
            if isinstance(ch, NavigableString):
                txt = re.sub(r"\s+", " ", str(ch).replace("\xa0", " "))
                if txt.strip() or (txt and state["p"].runs):
                    _set_font(state["p"].add_run(txt if state["p"].runs else txt.lstrip()), size, bold, italic)
            elif ch.name == "br":
                new_par()
            elif ch.name == "img":
                m = re.search(r"/figura/([0-9a-f]+)", ch.get("src", ""))
                if m and img_loader:
                    try:
                        state["p"].add_run().add_picture(io.BytesIO(img_loader(m.group(1))), width=Mm(55))
                    except Exception:
                        pass
            else:
                b, i = _style_flags(ch, bold, italic)
                if ch.name in BLOCK and state["p"].runs:
                    new_par()
                walk(ch, b, i)

    walk(cell_el, *_style_flags(cell_el, False, False))
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_docx(html: str, font_size: int, header_text: str, img_loader=None) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(15)
    sec.top_margin = sec.bottom_margin = Mm(15)
    sec.header_distance = Mm(7)
    hp = sec.header.paragraphs[0]
    _set_font(hp.add_run(header_text), font_size, bold=True)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(font_size)

    soup = BeautifulSoup(html, "html.parser")
    rows = soup.find_all("tr")
    cols = max(sum(int(td.get("colspan", 1) or 1) for td in tr.find_all("td", recursive=False)) for tr in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r_idx, tr in enumerate(rows):
        col = 0
        for td in tr.find_all("td", recursive=False):
            span = int(td.get("colspan", 1) or 1)
            cell = table.cell(r_idx, col)
            if span > 1:
                cell = cell.merge(table.cell(r_idx, min(col + span - 1, cols - 1)))
            style = (td.get("style") or "").replace(" ", "").lower()
            if "background-color:#f2f2f2" in style:
                _shade(cell, "F2F2F2")
            _fill(td, cell, font_size, "text-align:center" in style, img_loader)
            col += span
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
